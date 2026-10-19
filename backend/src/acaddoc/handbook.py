"""Programme handbook: cover, programme structure, programme outcomes with the
course-PO matrix, TOC, then every valid course.

Only courses with no blocking validation errors go in. The rest are reported,
so one bad syllabus cannot silently ship in a handbook.
"""
from __future__ import annotations

import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docxcompose.composer import Composer

from . import render
from .outcomes import STRENGTH_LABELS, Framework, course_level, mapping_programme, strength_mark
from .render import render_docx
from .schema import Course
from .validate import Issue, blocking, validate

_ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8}


@dataclass
class HandbookResult:
    included: list[Course] = field(default_factory=list)
    excluded: dict[str, list[Issue]] = field(default_factory=dict)


def semester_of(c: Course, programme: str) -> str | None:
    for o in c.offerings:
        if programme in o.programmes:
            return o.semester
    return None


def select(courses: list[Course], programme: str) -> list[tuple[str, Course]]:
    rows = [(semester_of(c, programme), c) for c in courses]
    rows = [(s, c) for s, c in rows if s]
    return sorted(rows, key=lambda r: (_ROMAN[r[0]], r[1].course_code))


def _title(s: str) -> str:
    small = {"and", "of", "for", "the", "in", "to"}
    return " ".join(w.lower() if w.lower() in small and i else w.capitalize()
                    for i, w in enumerate(s.split()))


def _field(paragraph, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        r = paragraph.add_run()
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)


def _outcomes_section(doc: Document, programme: str, rows: list[tuple[str, Course]], fw: Framework) -> None:
    """Programme outcomes, then how strongly each course mapped to them serves each one."""
    doc.add_paragraph("PROGRAMME OUTCOMES", style="AD Section").paragraph_format.page_break_before = True
    if fw.source != "programme":
        doc.add_paragraph(f"{programme} has not set its own programme outcomes; the UGC LOCF graduate "
                          "attributes are used.", style="AD Note")
    for o in fw.outcomes:
        p = doc.add_paragraph(style="AD Code")
        p.add_run(f"{o.code}\t", style="ADStrong")
        if o.title:
            p.add_run(f"{o.title}: ", style="ADStrong")
        p.add_run(o.text)

    mapped = [c for _, c in rows if mapping_programme(c) == programme and any(o.po for o in c.outcomes)]
    others = [c for _, c in rows if mapping_programme(c) != programme]
    unmapped = [c for _, c in rows if mapping_programme(c) == programme and not any(o.po for o in c.outcomes)]
    if not mapped:
        return
    doc.add_paragraph("COURSE \u2013 PROGRAMME OUTCOME MATRIX", style="AD Section")
    cols = fw.codes
    t = doc.add_table(rows=1 + len(mapped), cols=1 + len(cols))
    t.style = "Table Grid"
    t.autofit = False
    first = 2.0
    rest = min(1.4, (render.TEXT_WIDTH_CM - first) / max(1, len(cols)))

    def put(cell, text, strong=False, width=rest):
        cell.width = Cm(width)
        para = cell.paragraphs[0]
        para.style = "AD Cell"
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run(text, style="ADStrong" if strong else None)

    put(t.rows[0].cells[0], "Course", True, first)
    for j, code in enumerate(cols, 1):
        put(t.rows[0].cells[j], code, True)
    for i, c in enumerate(mapped, 1):
        put(t.rows[i].cells[0], c.course_code, True, first)
        for j, code in enumerate(cols, 1):
            put(t.rows[i].cells[j], strength_mark(course_level(c, code)))
    doc.add_paragraph("Each cell is the rounded average of the course's CO links to that outcome. 3 = "
                      f"{STRENGTH_LABELS[3]}, 2 = {STRENGTH_LABELS[2]}, 1 = {STRENGTH_LABELS[1]}, "
                      "- = not addressed.", style="AD Note")
    gaps = [code for code in cols if all(course_level(c, code) is None for c in mapped)]
    if gaps:
        doc.add_paragraph("Not addressed by any course in this handbook: " + ", ".join(gaps) + ".", style="AD Body")
    if unmapped:
        doc.add_paragraph("Not mapped yet: " + ", ".join(c.course_code for c in unmapped) + ".", style="AD Body")
    if others:
        doc.add_paragraph("Mapped to another programme's outcomes, so not in the matrix: "
                          + ", ".join(f"{c.course_code} ({mapping_programme(c) or 'none'})" for c in others) + ".",
                          style="AD Body")


def _front_matter(programme: str, year: str, rows: list[tuple[str, Course]],
                  fw: Framework | None = None) -> Document:
    doc = Document(str(render.TEMPLATE))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    for _ in range(3):
        doc.add_paragraph("", style="AD Body")
    for text, size in (("CURRICULUM HANDBOOK", 24), (programme, 20), (f"Academic Year {year}", 14)):
        p = doc.add_paragraph(style="AD Title")
        p.paragraph_format.page_break_before = False
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.size = Pt(size)

    doc.add_paragraph("PROGRAMME STRUCTURE", style="AD Section").paragraph_format.page_break_before = True
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, ["Semester", "Code", "Course", "L", "T", "P", "C"]):
        cell.text = ""; r = cell.paragraphs[0].add_run(h); r.bold = True
    total_c: dict[str, int] = {}
    for sem, c in rows:
        cells = t.add_row().cells
        for cell, v in zip(cells, [sem, c.course_code, _title(c.course_title), c.hours.lecture or "-",
                                   c.hours.tutorial or "-", c.hours.practical or "-", c.hours.credits]):
            cell.text = str(v)
        total_c[sem] = total_c.get(sem, 0) + c.hours.credits
    for row in t.rows:
        for cell, w in zip(row.cells, [2.0, 2.2, 8.0, 1.0, 1.0, 1.0, 1.0]):
            cell.width = Cm(w)
            for p in cell.paragraphs:
                p.style = "AD Cell"
    doc.add_paragraph("", style="AD Body")
    doc.add_paragraph("Credits per semester: " + ", ".join(f"{s}: {n}" for s, n in total_c.items()), style="AD Body")

    if fw is not None:
        _outcomes_section(doc, programme, rows, fw)

    doc.add_paragraph("CONTENTS", style="AD Section").paragraph_format.page_break_before = True
    _field(doc.add_paragraph(style="AD Body"), 'TOC \\u \\h \\z')
    return doc


def build_handbook(courses: list[Course], programme: str, year: str, out: Path, *,
                   include_invalid: bool = False,
                   framework_for: Callable[[str], Framework | None] | None = None) -> HandbookResult:
    """`framework_for(programme code)`: its outcomes, for the matrix and each course's CO-PO table."""
    result = HandbookResult()
    chosen = []
    fw_of = framework_for or (lambda _code: None)

    def course_fw(c: Course) -> Framework | None:
        return fw_of(mapping_programme(c))

    for sem, c in select(courses, programme):
        fw = course_fw(c)
        errs = blocking(validate(c, fw.outcomes if fw else None))
        if errs and not include_invalid:
            result.excluded[c.course_code] = errs
        else:
            chosen.append((sem, c)); result.included.append(c)

    master = _front_matter(programme, year, chosen, fw_of(programme))
    comp = Composer(master)
    with tempfile.TemporaryDirectory() as tmp:
        for _, c in chosen:
            p = Path(tmp) / f"{c.course_code}.docx"
            render_docx(c, p, allow_errors=True, framework=course_fw(c))
            comp.append(Document(str(p)))
    # ask Word to refresh the TOC when the file is opened
    settings = master.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
    out.parent.mkdir(parents=True, exist_ok=True)
    comp.save(str(out))
    return result
